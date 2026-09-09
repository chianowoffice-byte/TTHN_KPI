"""
Đọc sheet '2_Danh muc CV' của 8 file Excel cá nhân (đã có định giá đầy đủ) và
sinh ra sql/03-seed-danh-muc-cv.sql — 1 câu INSERT duy nhất join theo ma_cbnv,
không cần biết trước UUID của employees.

Chạy: python scripts/generate_job_catalog_seed.py "<đường dẫn thư mục Phòng QLNB>"
"""
import sys, os, io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# file Excel -> mã CBNV (suy từ nội dung Nhóm NV trong catalog + tên trong Sheet 1 gốc)
FILE_TO_MACBNV = {
    "ANHNTN13 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx": ("00146999", "Nguyễn Thị Nguyệt Anh"),
    "ANHNV14 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx":  ("00157279", "Nguyễn Vân Anh"),
    "THUYHT12 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx": ("00157267", "Hoàng Thị Thủy"),
    "DUONGTT4 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx": ("00157280", "Trần Thùy Dương"),
    "GIANGTTT4 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx":("00157275", "Trần Thị Thu Giang"),
    "MAINTP6 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx":  ("00169422", "Nguyễn Thị Phương Mai"),
    "NGAPT23 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx":  ("00183881", "Phạm Thanh Ngà"),
    "QUANGNH3 Nen_tang_do_luong_nang_suat_va_cham_KPI_Quy.xlsx": ("00157761", "Nguyễn Hồng Quang"),
}

def esc(v):
    if v is None:
        return "NULL"
    s = str(v).strip()
    if s == "":
        return "NULL"
    return "'" + s.replace("'", "''") + "'"

def esc_num(v):
    if v is None or str(v).strip() == "":
        return "NULL"
    return str(float(v))

def main(folder):
    rows = []
    skipped = []
    for fname, (ma_cbnv, ho_ten) in FILE_TO_MACBNV.items():
        path = os.path.join(folder, fname)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if '2_Danh muc CV' not in wb.sheetnames:
            print(f"-- BỎ QUA {fname}: không có sheet '2_Danh muc CV'")
            continue
        ws = wb['2_Danh muc CV']
        n = 0
        for row in ws.iter_rows(min_row=4, max_col=9, values_only=True):
            ma_cv, nhom_nv, ten_cv, dinh_ky, tc_td, tc_cl, a, b, c = row
            if not ma_cv and not ten_cv:
                continue
            if a is None or b is None or c is None:
                skipped.append((fname, ma_cv, ten_cv))
                continue
            rows.append((ma_cbnv, ma_cv, nhom_nv, ten_cv, dinh_ky, tc_td, tc_cl, a, b, c))
            n += 1
        wb.close()
        print(f"-- {fname} ({ho_ten}, {ma_cbnv}): {n} đầu việc")

    out_path = os.path.join(os.path.dirname(folder if False else __file__), '..', 'sql', '03-seed-danh-muc-cv.sql')
    out_path = os.path.abspath(out_path)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write("-- ============================================================================\n")
        f.write("-- SEED DANH MỤC & ĐỊNH GIÁ CÔNG VIỆC (job_catalog)\n")
        f.write(f"-- Sinh tự động từ sheet '2_Danh muc CV' của 8 file cá nhân — {len(rows)} đầu việc.\n")
        f.write("-- Chạy sau 02-seed-nhan-su.sql. Chạy lại an toàn (on conflict do nothing).\n")
        f.write("-- ============================================================================\n\n")
        f.write("insert into job_catalog\n")
        f.write("  (employee_id, ma_cv, nhom_nv, ten_cong_viec, dinh_ky_tan_suat, tieu_chi_tien_do, tieu_chi_chat_luong, pham_vi_anh_huong, muc_do_phuc_tap, thoi_gian_thuc_hien)\n")
        f.write("select e.id, v.ma_cv, v.nhom_nv, v.ten_cong_viec, v.dinh_ky, v.tc_td, v.tc_cl, v.a, v.b, v.c\n")
        f.write("from (values\n")
        lines = []
        for (ma_cbnv, ma_cv, nhom_nv, ten_cv, dinh_ky, tc_td, tc_cl, a, b, c) in rows:
            lines.append(
                f"  ({esc(ma_cbnv)},{esc(ma_cv)},{esc(nhom_nv)},{esc(ten_cv)},{esc(dinh_ky)},{esc(tc_td)},{esc(tc_cl)},{esc_num(a)},{esc_num(b)},{esc_num(c)})"
            )
        f.write(",\n".join(lines))
        f.write("\n) as v(ma_cbnv, ma_cv, nhom_nv, ten_cong_viec, dinh_ky, tc_td, tc_cl, a, b, c)\n")
        f.write("join employees e on e.ma_cbnv = v.ma_cbnv\n")
        f.write("on conflict (employee_id, ma_cv) do nothing;\n")

    print(f"\n-- Đã ghi {out_path} — {len(rows)} dòng.")
    if skipped:
        print(f"-- Bỏ qua {len(skipped)} dòng thiếu (a)(b)(c):")
        for s in skipped[:20]:
            print(f"--   {s}")

if __name__ == "__main__":
    main(sys.argv[1])
